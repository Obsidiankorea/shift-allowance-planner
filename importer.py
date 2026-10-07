"""Importer for the supplied Gyeongnam shift workbook; does not execute macros/formulas."""
import datetime as dt, math, re, sys
# openpyxl optionally imports numpy. The EXE bundles none, so a stray/broken numpy found on another PC
# crashed start-up ("module 'numpy' has no attribute 'short'"). Block it in the frozen build only.
if getattr(sys,'frozen',False):sys.modules.setdefault('numpy',None)
import openpyxl
from domain import *

def normalize(v):
    if v is None:return '비'
    s=re.sub(r'\s+','',str(v))
    return {'주간':'주','야간':'야','비번':'비','휴무':'비','주간연가':'주연','야간연가':'야연','주간특가':'주특','야간특가':'야특','주간공가':'주공','야간공가':'야공','관외출장':'출장','집합교육':'교육'}.get(s,s)

def import_excel(path):
    w=openpyxl.load_workbook(path,data_only=True)
    required=['월중 근무내역','초과근무수당산출내역서','기준표']
    if any(n not in w.sheetnames for n in required):raise ValueError('지원 서식: 월중 근무내역 / 초과근무수당산출내역서 / 기준표 시트가 필요합니다.')
    p=empty_project();p['source']=str(path);s=w['월중 근무내역'];summary=w[required[1]]
    cols={c:s.cell(2,c).value.date() for c in range(8,s.max_column+1) if isinstance(s.cell(2,c).value,dt.datetime)}
    if not cols:raise ValueError('2행의 실제 날짜를 찾을 수 없습니다. Excel에서 재계산 후 저장해 주세요.')
    first=min(cols.values());month=first.strftime('%Y-%m');p['import_month']=month;p['rate_year']=first.year
    for row in w['기준표'].iter_rows():
        if row[0].value in RATES and all(isinstance(x.value,(int,float)) for x in row[1:4]):p['rates'][row[0].value]=[x.value for x in row[1:4]]
    headings={str(s.cell(2,c).value).replace('\n','').replace(' ',''):c for c in range(1,s.max_column+1)}
    def num(r,label,default=0):
        c=headings.get(label)
        v=s.cell(r,c).value if c else default
        return float(v or 0)
    activities={}
    if '비번활동 내역' in w.sheetnames:
        for row in w['비번활동 내역'].iter_rows(min_row=9):
            name,date,h=row[1].value,row[2].value,row[8].value
            if name and isinstance(date,dt.datetime) and isinstance(h,(int,float)):
                key=(str(name).strip(),str(date.date()));activities[key]=activities.get(key,0)+h
    for r in range(3,s.max_row+1):
        name=s.cell(r,4).value;grade=s.cell(r,3).value;team=s.cell(r,2).value
        if not isinstance(name,str) or grade not in p['rates'] or not team:continue
        pid=f'p{r}';person={'id':pid,'name':name.strip(),'grade':grade,'team':str(team),'salary':0}
        p['people'].append(person);p['baseline'][pid]={};numeric=0;activity=0
        for c,d in cols.items():
            v=s.cell(r,c).value;e={'code':normalize(v)}
            if isinstance(v,(int,float)):e={'code':'비','extra':float(v)};numeric+=v
            if e['code'] not in CODES:p['warnings'].append(f'{name} {d}: 미등록 코드 {e["code"]} — 날짜를 편집해 주세요.')
            a=activities.get((name.strip(),str(d)),0)
            if a:e['activity']=a;activity+=a
            p['baseline'][pid][str(d)]=e
        ap=num(r,'기타시간합계');ar=num(r,'야간제외시간')
        p['adjustments'][pid]={month:{'hours':ap-numeric-math.floor(activity),'night':-ar,'basic':float(summary.cell(r-1,13).value or 0)}}
        p['reference'][pid]={k:summary.cell(r-1,c).value for k,c in [('total',14),('base_days',15),('signed',17),('night',18),('pay',20),('overtime',21),('night_pay',22),('holiday_pay',23)]}
        if p['reference'][pid]['pay'] is None:raise ValueError('수식의 저장된 계산 결과가 없습니다. Excel에서 계산 후 저장한 파일을 선택해 주세요.')
    if not p['people']:raise ValueError('지원되는 직급의 인원을 찾을 수 없습니다.')
    for team in sorted({x['team'] for x in p['people']}):
        records=[(dt.date.fromisoformat(d),e['code']) for x in p['people'] if x['team']==team for d,e in p['baseline'][x['id']].items() if e['code'] in ['주','야','비']]
        scores=[sum(CYCLE[((d-first).days+phase)%4]!=code for d,code in records) for phase in range(4)]
        p['teams'][team]={'anchor':str(first),'phase':scores.index(min(scores))}
    p['source_holidays']={}
    if '휴일' in w.sheetnames:
        for row in w['휴일']:
            for cell in row:
                if isinstance(cell.value,dt.datetime):p['source_holidays'][str(cell.value.date())]='엑셀 휴일'
    for person in p['people']:
        calc=calculate(p,person['id'],first.year,first.month)
        diffs=[k for k,v in p['reference'][person['id']].items() if v is not None and abs(calc[k]-v)>.001]
        if diffs:p['warnings'].append(f'{person["name"]}: 엑셀과 불일치({", ".join(diffs)}). 휴일/코드/월 조정값 확인 필요.')
    return p

legacy_import=import_excel

def import_excel(path):
    w=openpyxl.load_workbook(path,data_only=True)
    if '근무편성표' not in w.sheetnames:return legacy_import(path)
    from collections import Counter
    p=empty_project();p['source']=str(path)
    roster=next((s for s in w if '근무자 입력' in s.title),None)
    if roster is None:raise ValueError('근무자 입력 시트를 찾지 못했습니다.')
    clean=lambda x:re.sub(r'\s+','',str(x or ''))
    def grade(v):
        v=clean(v)
        if v.startswith('소방'):return v
        m=re.search(r'(\d+)급',v)
        if not m:raise ValueError(f'직급을 해석할 수 없습니다: {v}')
        return '일반'+m.group(1)
    byname={}
    for r in range(3,roster.max_row+1):
        team=clean(roster.cell(r,5).value)
        if not re.fullmatch(r'\d+팀',team):continue
        for g,n in [(6,7),(8,9),(10,11)]:
            name=clean(roster.cell(r,n).value)
            if not name:continue
            if name in byname:raise ValueError(f'동명이인/중복 이름 확인 필요: {name}')
            pid=f'person{len(byname)+1}';person={'id':pid,'name':name,'grade':grade(roster.cell(r,g).value),'team':team,'salary':0}
            byname[name]=person;p['people'].append(person)
    s=w['근무편성표'];rows=[];date=None;team_records={};notes=[]
    for r in range(5,s.max_row+1):
        v=s.cell(r,1).value
        if isinstance(v,dt.datetime):date=v.date()
        code=normalize(s.cell(r,3).value)
        if date is None or code not in ['주','야','당']:continue
        team=clean(s.cell(r,4).value)
        team_records.setdefault(team,[]).append((date,code))
        names=[clean(s.cell(r,c).value) for c in [6,8,10]]
        rows.append((date,code,team,names));note=clean(s.cell(r,11).value)
        if note:notes.append((date,code,note))
    if not rows:raise ValueError('근무편성표에 저장된 날짜/근무 계산값이 없습니다. Excel 재계산 후 저장해 주세요.')
    first=min(r[0] for r in rows);p['import_month']=first.strftime('%Y-%m')
    for team,records in team_records.items():
        scores=[sum(CYCLE[((d-first).days+phase)%4]!=c for d,c in records) for phase in range(4)]
        p['teams'][team]={'anchor':str(first),'phase':scores.index(min(scores))}
    for person in p['people']:
        p['baseline'][person['id']]={str(d):{'code':'비'} for d in dates(first.year,first.month)}
    hs=holiday_map(p,first.year)
    for d,c,team,names in rows:
        for name in names:
            if not name:continue
            if name not in byname:raise ValueError(f'근무자 입력 시트에 없는 이름: {name}. 명부를 먼저 수정해 주세요.')
            old=p['baseline'][byname[name]['id']][str(d)]
            if old['code']=='비':e={'code':c}
            else:
                e={'code':'복합','credit':old.get('credit',WORK.get(old['code'],0))+WORK[c], 'night':old.get('night',8 if old['code'] in ['야','당'] else 0)+(8 if c in ['야','당'] else 0),'holiday':old.get('holiday',int(old['code'] in ['주','당'] and (d.weekday()>=5 or str(d) in hs)))+int(c in ['주','당'] and (d.weekday()>=5 or str(d) in hs)), 'memo':'동일 날짜의 주간·야간 근무 합산'}
                p['warnings'].append(f'{name} {d}: 같은 날짜에 복수 근무를 합산했습니다. 실제 편성 확인 필요.')
            p['baseline'][byname[name]['id']][str(d)]=e
    for d,c,note in notes:
        matched=False
        for name,person in byname.items():
            if name not in note:continue
            m=re.search(re.escape(name)+r'(연가|교육|출장|공가|특가|병가)',note)
            if not m:continue
            kind=m.group(1);code={'연가':c+'연','공가':c+'공','특가':c+'특','병가':c+'병','교육':'교육','출장':'출장'}[kind]
            old=p['baseline'][person['id']][str(d)]
            if old['code']!='비':
                p['warnings'].append(f'{name} {d}: 근무와 비고({note})가 중복됩니다. 수동 확인 필요.');continue
            e={'code':code,'memo':'엑셀 비고: '+note}
            if kind in ['교육','출장']:e['hours']=8;e['memo']+=' / 시간 미기재: 임시 8시간, 확인 필요';p['warnings'].append(f'{name} {d} {kind}: 비고에 시간이 없어 8시간을 임시 반영했습니다. 수정해 주세요.')
            p['baseline'][person['id']][str(d)]=e;matched=True
        if not matched:p['warnings'].append(f'{d} 미해석 비고: {note}')
    p['source_holidays']={str(roster.cell(r,1).value.date()):str(roster.cell(r,2).value or '엑셀 휴일') for r in range(3,roster.max_row+1) if isinstance(roster.cell(r,1).value,dt.datetime)}
    p['warnings'].append('편성표에는 단가표가 없어 제공된 2026년 수당 산출표의 단가를 적용했습니다. 설정에서 변경할 수 있습니다.')
    return p
